import os
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import google.generativeai as genai
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Configure the API key for Google Generative AI
genai.configure(api_key="AIzaSyC7W9OoTb7cZ_top7bFMK10sLGFkjeS5zA")
model = genai.GenerativeModel("gemini-1.5-flash")

def add_custom_paragraph(doc, text, left_indent):
    """
    Add a paragraph with custom formatting based on provided parameters.
    """
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True  # Make text bold

    # Set indentation
    paragraph_format = paragraph.paragraph_format
    paragraph_format.left_indent = Pt(left_indent)

def process_plain_text_to_word(plain_text, output_file="structured_document.docx"):
    """
    Process plain text into a structured Word document with proper formatting.
    """
    try:
        print(f"Received text for processing:\n{plain_text}")

        # Prepare the prompt for the Generative AI
        prompt = (
            "Analyze the following plain text and organize it into a structured document with appropriate headings like proper headings,sub headings,sub-sub headings,good paragraphs\n"
            "Follow the given five rules strictly all and generate a document with the structured content.\n"
            "rule1: Actual purpose of the report or the Actual Title should start with % and ends with %\n"
            "rule2: Heading should start with # and ends with # and also maintain indexing like 1,2,3,..\n"
            "rule3: Subheading should start with $ and ends with $ and also maintain indexing like 1.1,1.2,1.3,..\n"
            "rule4: Sub-Subheading should start with * and ends with * and also maintain indexing like a,b,c..\n"
            "rule5: Paragraph should start with -- and ends with --\n"
            f"\n{plain_text}\n"
            "Output the structure in a hierarchical format (e.g., numbered headings, subheadings, and paragraphs)."
        )
        print(f"Prompt for the AI model:\n{prompt}")
        # Generate the structured text using the AI model
        response = model.generate_content(prompt)
        if not hasattr(response, 'text'):
            raise ValueError("Invalid response from the AI model")

        structured_text = response.text
        print(f"Structured response:\n{structured_text}")

        # Create a Word document
        doc = Document()
            # Set page margins (optional - you can adjust these values)
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.5)
            section.right_margin = Inches(0.5)
        
            # Get the section's properties
            section_properties = section._sectPr
        
            # Create border elements for all sides
            borders_xml = parse_xml(f'''
                <w:pgBorders {nsdecls('w')}>
                    <w:top w:val="single" w:sz="24" w:space="0" w:color="000000"/>
                    <w:left w:val="single" w:sz="24" w:space="0" w:color="000000"/>
                    <w:bottom w:val="single" w:sz="24" w:space="0" w:color="000000"/>
                    <w:right w:val="single" w:sz="24" w:space="0" w:color="000000"/>
                </w:pgBorders>
            ''')
        
            # Add borders
            section_properties.append(borders_xml)
            
        lines = structured_text.split("\n")
        for line in lines:
            if line.startswith("%") and line.endswith("%"):
                content = line.strip('%')
                add_custom_paragraph(doc, content, left_indent=3)
            elif line.startswith("#") and line.endswith("#"):
                content = line.strip('#')
                add_custom_paragraph(doc, content, left_indent=9)
            elif line.startswith("$") and line.endswith("$"):
                content = line.strip('$')
                add_custom_paragraph(doc, content, left_indent=12)
            elif line.startswith("*") and line.endswith("*"):
                content = line.strip('*')
                add_custom_paragraph(doc, content, left_indent=18)
            else:
                # Regular paragraph
                paragraph = doc.add_paragraph(line)
                paragraph.paragraph_format.left_indent = Pt(36)  # Default paragraph indentation

        # Save the document
        doc.save(output_file)
        print(f"Document saved as {output_file}")
        return output_file

    except Exception as e:
        print(f"Error: {str(e)}")
        return None

